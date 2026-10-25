import io
import sys
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.writer.excel import ExcelWriter
from openpyxl.xml import LXML

HEADER = ['Item', 'Dia', 'Valor', 'Quem', 'Pago?']


def household_month(title, income, fixed, variable):
    blocks = [
        ('RECEITAS', 'Total receitas', income),
        ('FIXAS', 'Total fixas', fixed),
        ('VARIÁVEIS', 'Total variáveis', variable),
    ]
    rows = [[title], []]
    balance = 0
    for name, label, lines in blocks:
        total = round(sum(line[2] for line in lines), 2)
        balance += total if name == 'RECEITAS' else -total
        rows += [[name], HEADER, *lines, [label, None, total], []]
    return [*rows, ['SALDO DO MÊS', None, round(balance, 2)]]


HOUSEHOLD = {
    'Set-26': household_month(
        'Orçamento da casa - setembro/2026',
        [['Salário Marcio', 5, 6800, 'Marcio', 'sim'], ['Salário Gabriela', 5, 5900, 'Gabi', 'sim']],
        [
            ['Aluguel', 10, 2200, 'Conjunto', 'sim'],
            ['Condomínio', 10, 650, 'Conjunto', 'sim'],
            ['Internet', 15, 120, 'Conjunto', 'sim'],
            ['Parcela carro (2/36)', 20, 980, 'Marcio', 'sim'],
        ],
        [['Mercado', 8, 1210.3, 'Conjunto', 'sim'], ['Farmácia', 12, 64.5, 'Gabi', 'sim']],
    ),
    'Out-26': household_month(
        'Orçamento da casa - outubro/2026',
        [
            ['Salário Marcio', 5, 6800, 'Marcio', 'sim'],
            ['Salário Gabriela', 5, 5900, 'Gabi', 'sim'],
            ['Freela site', 20, 1200, 'Marcio', 'não'],
        ],
        [
            ['Aluguel', 10, 2200, 'Conjunto', 'sim'],
            ['Condomínio', 10, 650, 'Conjunto', 'sim'],
            ['Internet', 15, 120, 'Conjunto', 'sim'],
            ['Parcela carro (3/36)', 20, 980, 'Marcio', 'não'],
        ],
        [
            ['Mercado', 8, 1340.5, 'Conjunto', 'sim'],
            ['Farmácia', 12, 89.9, 'Gabi', 'sim'],
            ['Gasolina', 14, 310, 'Marcio', 'sim'],
            ['Restaurante', 18, 220, 'Conjunto', 'não'],
        ],
    ),
}

PLANNING = {
    'Nov-26': [
        ['Planejamento de novembro de 2026'],
        [],
        ['Receitas', None, None, None, 'Fixas', None, None, None, 'Variáveis'],
        ['Descrição', 'Dia', 'Valor', None, 'Descrição', 'Dia', 'Valor', None, 'Descrição', 'Quem', 'Valor'],
        ['Salário Marcio', 5, 6800, None, 'Aluguel', 10, 2200, None, 'Mercado', 'Conjunto', 1400],
        ['Salário Gabriela', 5, 5900, None, 'Condomínio', 10, 650, None, 'Farmácia', 'Conjunto', 150],
        ['Total', None, 12700, None, 'Internet', 15, 120, None, 'Lazer', 'Conjunto', 300],
        [None, None, None, None, 'Luz', 18, 210, None, 'Presentes', 'Gabi', 200],
        [None, None, None, None, 'Total', None, 3180, None, 'Total', None, 2050],
        [],
        [],
        ['Saldo previsto', None, 7470],
    ],
}

STATEMENT = [
    [
        'Extrato de conta corrente - outubro de 2026 - página 1 de 2',
        'Titular: Gabriela - Conta Gabriela, final 4321',
        ('Data   Descrição', 'Valor (R$)'),
        ('01/10   Saldo anterior', '3.250,00'),
        ('05/10   Salário Empresa Fictícia', '5.900,00'),
        ('07/10   Academia Bem-Estar', '-129,90'),
        ('10/10   Farmácia Vida Fictícia', '-86,40'),
        ('15/10   Curso de inglês online', '-350,00'),
    ],
    [
        'Extrato de conta corrente - outubro de 2026 - página 2 de 2',
        ('Data   Descrição', 'Valor (R$)'),
        ('18/10   Restaurante Sabor Fictício', '-142,70'),
        ('22/10   Notebook parcela 4/10', '-289,00'),
        ('25/10   Transporte por aplicativo', '-47,80'),
        ('Total de créditos', '5.900,00'),
        ('Total de débitos', '1.045,80'),
        ('Saldo final em 31/10', '8.104,20'),
    ],
]


def workbook(sheets, created, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    book.properties.created = created
    book.properties.modified = created
    written = io.BytesIO()
    with zipfile.ZipFile(written, 'w', zipfile.ZIP_DEFLATED) as archive:
        ExcelWriter(book, archive).save()
    with zipfile.ZipFile(written) as source, zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            entry = zipfile.ZipInfo(item.filename, created.timetuple()[:6])
            entry.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(entry, source.read(item))


def statement(pages, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = FPDF(format='A4')
    pdf.set_creation_date(datetime(2026, 10, 31, tzinfo=UTC))
    pdf.set_producer('')
    pdf.set_auto_page_break(False)
    pdf.set_font('Helvetica', size=11)
    for lines in pages:
        pdf.add_page()
        for line in ['Banco Fictício S.A.', *lines]:
            text, value = line if isinstance(line, tuple) else (line, '')
            pdf.cell(130, 8, text)
            pdf.cell(40, 8, value, align='R', new_x='LMARGIN', new_y='NEXT')
    pdf.output(str(path))


def main(root):
    if not LXML:
        sys.exit('fixtures need lxml: run them through npm run eval:fixtures')
    cases = Path(root)
    workbook(HOUSEHOLD, datetime(2026, 10, 25), cases / 'household-xlsx' / 'orcamento-casa.xlsx')
    workbook(PLANNING, datetime(2026, 10, 25), cases / 'planning-xlsx' / 'planejamento-lado-a-lado.xlsx')
    statement(STATEMENT, cases / 'statement-pdf' / 'extrato-gabriela.pdf')


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).parents[1] / 'cases')
