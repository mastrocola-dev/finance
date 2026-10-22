from pathlib import Path

from openpyxl import load_workbook

from ingest.segment import Segment, plain
from ingest.table import table_segments


def read_xlsx(path: Path) -> list[Segment]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            segment
            for sheet in workbook.worksheets
            for segment in table_segments(
                sheet.title, [[plain(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            )
        ]
    finally:
        workbook.close()
