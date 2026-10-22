from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

FIXTURES = Path(__file__).parent / 'fixtures'


def write_pdf(path: Path, pages: list[list[str]]) -> Path:
    count = len(pages)
    objects = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids ['
        + b' '.join(f'{4 + 2 * index} 0 R'.encode() for index in range(count))
        + f'] /Count {count} >>'.encode(),
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>',
    ]
    for index, lines in enumerate(pages):
        text = ' T* '.join(f'({line})Tj' for line in lines)
        stream = f'BT /F1 11 Tf 14 TL 72 760 Td {text} ET'.encode('cp1252')
        objects.append(
            f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> '
            f'/Contents {5 + 2 * index} 0 R >>'.encode()
        )
        objects.append(f'<< /Length {len(stream)} >>\nstream\n'.encode() + stream + b'\nendstream')
    body = b'%PDF-1.4\n'
    offsets = []
    for number, content in enumerate(objects, 1):
        offsets.append(len(body))
        body += f'{number} 0 obj\n'.encode() + content + b'\nendobj\n'
    xref = len(body)
    body += f'xref\n0 {len(objects) + 1}\n0000000000 65535 f \n'.encode()
    body += b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets)
    body += f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    path.write_bytes(body)
    return path


@pytest.fixture
def fixtures() -> Path:
    return FIXTURES


@pytest.fixture
def household(tmp_path: Path) -> Path:
    workbook = Workbook()
    october = workbook.active
    october.title = '2026-10'
    october['A1'] = 'Orçamento de outubro'
    rows = {
        3: ['Receitas', 'Quem', 'Valor'],
        4: ['Salário', 'M', 5000],
        5: ['Freelance', 'G', 1200.5],
        8: ['Despesas fixas', 'Vencimento', 'Valor', 'Pago'],
        9: ['Aluguel', date(2026, 10, 5), 2400, True],
        10: ['Internet', date(2026, 10, 10), 119.9, False],
    }
    for row, values in rows.items():
        for column, value in enumerate(values, 1):
            october.cell(row, column, value)
    october['F3'] = 'Total receitas'
    october['G3'] = 6200.5
    october['F4'] = 'Total despesas'
    october['G4'] = 2519.9
    notes = workbook.create_sheet('Notas da casa')
    notes['B2'] = 'Planilha fictícia'
    path = tmp_path / 'casa.xlsx'
    workbook.save(path)
    return path
